"""
替换开题报告中的项目简介，扩展到约 300 字并与最新方案对齐。
"""

import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

doc = Document('files/面向电子电路学习的图像识别和可解释分析辅助系统-开题报告.docx')

# TABLE 1 项目简介
table1 = doc.tables[1]
cell = table1.rows[0].cells[0]

# 清除所有现有 run
for p in cell.paragraphs:
    for r in p.runs:
        r._element.getparent().remove(r._element)

# 获取第一个段落
p = cell.paragraphs[0]


def make_run(paragraph, text, font_name='楷体', font_size_pt=11):
    run = paragraph.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size_pt)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = run._element.makeelement(qn('w:rFonts'), {})
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:eastAsia'), font_name)
    return run


# 写入新的项目简介（约 300 字）
make_run(p, '本项目拟结合图像处理、类型化电路拓扑建模、符号计算和教学可视化方法，')
make_run(p, '研发一套面向课程学习的电路智能分析与教学辅助系统。')
make_run(p, '系统采用 YOLOv8 + PaddleOCR + OpenCV 实现电路图自动识别，')
make_run(p, '采用 Zod Schema 定义类型化电路图数据结构，')
make_run(p, '采用 Lcapy + SymPy 实现符号电路分析，')
make_run(p, '采用 React Flow + KaTeX 实现电路图编辑和公式渲染。')
make_run(p, '系统以电脑端本地软件为主体，支持手机局域网扫码拍照上传电路图；')
make_run(p, '电脑端完成图像处理、候选拓扑生成、用户校正、自动分析和结果展示，')
make_run(p, '输出等效电路、增益表达式、分步骤解题过程和知识点标注。')
make_run(p, '项目既服务于学生完成习题和理解等效模型，')
make_run(p, '也具有电路图理解和智能教育工具构建方面的研究价值。')

# 保存
doc.save('files/面向电子电路学习的图像识别和可解释分析辅助系统-开题报告.docx')
print(f'项目简介已更新，总字数: {len(cell.text)}')
