"""
替换开题报告中数据层、分析引擎层、展示层相关内容。
保持原有字体（楷体）、字号（12pt）、行间距（固定值20pt）不变。
"""

import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

doc = Document('files/面向电子电路学习的图像识别和可解释分析辅助系统-开题报告.docx')


def make_run(paragraph, text, font_name='楷体', font_size_pt=12, bold=False):
    run = paragraph.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = run._element.makeelement(qn('w:rFonts'), {})
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:eastAsia'), font_name)
    return run


def clear_paragraph_runs(paragraph):
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)


# ============================================================
# TABLE 3 研究方案
# ============================================================
table3 = doc.tables[3]
cell = table3.rows[0].cells[0]

# 段落5: 电路分析方法
p5 = cell.paragraphs[5]
clear_paragraph_runs(p5)
make_run(p5, '（2）电路分析方法：')
make_run(p5, '采用"课程规则层 + 符号/仿真工具层"的双层架构。')
make_run(p5, '课程规则层负责题型分类（二极管/BJT/MOSFET 共 5 类策略）、分析策略选择、等效电路生成和分步骤解题过程；')
make_run(p5, '符号/仿真工具层采用 Lcapy（MNA 符号求解）+ SymPy（表达式化简/Laplace 变换）+ PySpice（数值校验）的工具链，')
make_run(p5, '通过插件式分析流水线（Strategy Pattern）支持题型扩展。')
make_run(p5, '求解器采用三重保护机制：超时控制（10 秒阈值）→ 复杂度限制（表达式项数 >100 时化简）→ 降级为 PySpice 数值求解，防止符号计算卡死。')

# 段落6: 数据建模方法
p6 = cell.paragraphs[6]
clear_paragraph_runs(p6)
make_run(p6, '（3）数据建模方法：')
make_run(p6, '采用 TypeScript + Zod Schema 作为数据模型的事实来源（Schema-as-Source-of-Truth），')
make_run(p6, '通过 zod-to-json-schema 自动导出 JSON Schema 供 Python 端（Pydantic）消费，实现 TypeScript ↔ Python 跨语言互操作。')
make_run(p6, '电路图采用 NEC（节点-元件-连接）三元组表示范式，与 KiCad 等工业 EDA 工具的数据模型对齐。')
make_run(p6, '元件定义采用五维结构（端口/参数/模型变体/分析标签/教学标签），覆盖 19 个元件族。')
make_run(p6, '校验器采用"报告不阻断"设计，返回 ValidationIssue[] 而非抛出异常，适配半自动工作流。')

# 段落10: 第三阶段
p10 = cell.paragraphs[10]
clear_paragraph_runs(p10)
make_run(p10, '第三阶段（2026年12月-2027年1月）：完成分析引擎层构建，包括插件式分析流水线（5 类题型策略）、')
make_run(p10, 'MNA 矩阵建模与 Lcapy 符号求解、三重保护机制（超时/复杂度限制/降级）、')
make_run(p10, 'SymPy/Lcapy/SPICE 数值校验和结构化 AnalysisResult 输出。')

# 段落11: 第四阶段
p11 = cell.paragraphs[11]
clear_paragraph_runs(p11)
make_run(p11, '第四阶段（2027年2-4月）：完成展示层制作，包括 React Flow 电路图编辑器（自定义元件节点+拖拽连线）、')
make_run(p11, 'KaTeX 公式渲染（~10ms/公式）、等效电路图展示、分步骤解题交互、')
make_run(p11, '报告导出（PDF/Word/LaTeX 三种格式）、QR 码 + SSE 局域网手机拍照上传联动。')

# ============================================================
# TABLE 4 预期成果
# ============================================================
table4 = doc.tables[4]
cell4 = table4.rows[0].cells[0]

# 段落2: 数据层
p2 = cell4.paragraphs[2]
clear_paragraph_runs(p2)
make_run(p2, '（1）数据层：TypeScript + Zod Schema 作为事实来源，NEC（节点-元件-连接）三元组电路图表示，')
make_run(p2, '19 个元件的五维定义（端口/参数/模型变体/分析标签/教学标签），')
make_run(p2, '3 个 Benchmark 样例（二极管/BJT/MOS），JSON Schema 跨语言导出，')
make_run(p2, '"报告不阻断"校验器，SQLite 任务队列。')

# 段落4: 分析引擎层
p4 = cell4.paragraphs[4]
clear_paragraph_runs(p4)
make_run(p4, '（3）分析引擎层：双层架构（课程规则层 + 符号/仿真工具层），')
make_run(p4, '插件式分析流水线（Strategy Pattern，5 类题型策略），')
make_run(p4, 'Lcapy + SymPy + PySpice 工具链，MNA 矩阵建模，')
make_run(p4, '三重保护机制（超时/复杂度限制/降级为数值求解），')
make_run(p4, '输出结构化 AnalysisResult JSON（等效电路、关键方程、分步骤推导、知识点标注）。')

# 段落5: 展示层
p5r = cell4.paragraphs[5]
clear_paragraph_runs(p5r)
make_run(p5r, '（4）展示层：React Flow 电路图编辑器（自定义元件节点、拖拽连线、参数编辑），')
make_run(p5r, 'KaTeX 公式渲染（~10ms/公式），等效电路图展示，分步骤解题交互（展开/折叠/跳转），')
make_run(p5r, '报告导出（PDF/Word/LaTeX 三种格式），QR 码 + SSE 局域网手机拍照上传联动。')

# 保存
doc.save('files/面向电子电路学习的图像识别和可解释分析辅助系统-开题报告.docx')
print('DOCX 更新完成。')
