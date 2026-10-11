from docx import Document
from docx.shared import Pt, Twips, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# Helper function to set font for a run
def set_run_font(run, font_name, font_size_half_points, bold=False):
    run.font.name = font_name
    run.font.size = Pt(font_size_half_points / 2)  # Convert half-points to points
    if bold:
        run.font.bold = True
    # Set east asian font
    rPr = run.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)

# Helper function to add section heading (B style)
def add_section_heading(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    # Section headings use bold
    run.font.bold = True
    run.font.size = Pt(12)  # Default size
    set_run_font(run, '宋体', 24)
    return p

# Helper function to add subsection heading (华文中宋, 24半角)
def add_subsection_heading(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, '华文中宋', 24)
    p.paragraph_format.first_line_indent = Twips(120)
    p.paragraph_format.space_before = Twips(156)
    return p

# Helper function to add body paragraph (楷体, 24半角 or 22半角)
def add_body_paragraph(text, font_size=24, first_line_indent=480, line_spacing=400):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run, '楷体', font_size)
    p.paragraph_format.first_line_indent = Twips(first_line_indent)
    p.paragraph_format.line_spacing = Twips(line_spacing)
    return p

# Helper function to add numbered item (楷体, 24半角)
def add_numbered_item(text, font_size=24):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run, '楷体', font_size)
    p.paragraph_format.first_line_indent = Twips(480)
    p.paragraph_format.line_spacing = Twips(400)
    return p

# ===== 四、项目实施方案 =====
add_section_heading('四、项目实施方案')

# （一）研究方案
add_subsection_heading('（一）研究方案')

add_numbered_item('1、理论基础与技术路线')
add_body_paragraph('本项目以图像处理、类型化电路拓扑建模、符号计算和教学可视化为理论基础，采用"数据层 → 图像处理层 → 分析引擎层 → 展示层"的分层架构设计。系统以电脑端本地软件为主体，支持手机在同一局域网内扫码拍照上传电路图，实现从图像采集到教学级分析报告的完整闭环。')

add_numbered_item('2、研究方法')
add_body_paragraph('（1）图像处理方法：采用 OpenCV 进行预处理（灰度化、二值化、去噪、透视/倾斜校正），结合元件检测和 OCR 参数提取，生成带置信度的候选拓扑。对于低置信度元件、疑似断线、未连接端口等问题，系统将显式提示给用户确认。')
add_body_paragraph('（2）电路分析方法：采用"课程规则层 + 符号/仿真工具层"的混合方案。课程规则负责等效变换、器件模型和题型逻辑，SymPy、Lcapy、SPICE 等工具用于符号求解、线性电路分析和数值校验，通过插件式分析流程和分题型策略实现可测试、可追溯的分析过程。')
add_body_paragraph('（3）数据建模方法：采用 TypeScript / Zod Schema 定义类型化电路图数据结构，建立元件库和 Benchmark 数据集，支持 SPICE / netlist 导出和校验，确保数据的一致性和可验证性。')

add_numbered_item('3、研究步骤')
add_body_paragraph('第一阶段（2026年6-8月）：完成数据层建设，包括类型化电路图 Schema 定义、元件库建立、Benchmark 样例收集与采集规范制定。')
add_body_paragraph('第二阶段（2026年9-12月）：完成图像处理层开发，包括 OpenCV 预处理、元件/OCR 识别、候选拓扑生成与置信度标注。')
add_body_paragraph('第三阶段（2026年12月-2027年1月）：完成分析引擎层构建，包括题型策略、课程规则、SymPy/Lcapy/SPICE 校验和 LaTeX 推导输出。')
add_body_paragraph('第四阶段（2027年2-4月）：完成展示层制作，包括电路编辑器、等效电路展示、步骤解释、报告导出和手机-电脑联动体验。')

add_numbered_item('4、人员分工')
add_body_paragraph('项目负责人王子涵（24211119）：负责整体架构设计、数据层开发和项目管理。')
add_body_paragraph('项目成员金浩轩（24211085）：负责图像处理层开发、元件识别和 OCR 参数提取。')
add_body_paragraph('项目成员刘勃源（24211088）：负责分析引擎层开发、课程规则和符号计算集成。')

add_numbered_item('5、计划使用的学校资源条件')
add_body_paragraph('（1）实验室资源：使用电子信息工程学院的计算机实验室进行开发和测试。')
add_body_paragraph('（2）软件资源：使用学校提供的正版软件授权，包括 MATLAB、Multisim 等仿真工具。')
add_body_paragraph('（3）数据资源：收集典型教材题和习题册，建立开发测试集和评测基准。')

# （二）实施进度安排
add_subsection_heading('（二）实施进度安排')

# Create schedule table
table = doc.add_table(rows=5, cols=4)
table.style = 'Table Grid'

# Set table headers
headers = ['序号', '项目实施各阶段内容', '时间安排', '备注']
for i, header in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = header
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_run_font(run, '华文中宋', 24)

# Table data
data = [
    ['1', '完成类型化电路图 Schema、元件库、benchmark 样例与采集规范，沉淀可复现数据标准', '2026.6-8', '数据层建设'],
    ['2', 'OpenCV 预处理、元件/OCR 识别、候选拓扑与置信度标注', '2026.9-12', '图像处理层开发'],
    ['3', '构建题型策略、课程规则、SymPy/Lcapy/SPICE 校验和 LaTeX 推导输出', '2026.12-2027.1', '分析引擎层构建'],
    ['4', '制作电路编辑器、等效电路展示、步骤解释、报告导出和手机-电脑联动体验', '2027.2-4', '展示层制作'],
]

for row_idx, row_data in enumerate(data):
    for col_idx, cell_text in enumerate(row_data):
        cell = table.rows[row_idx + 1].cells[col_idx]
        cell.text = cell_text
        for paragraph in cell.paragraphs:
            if col_idx == 0:  # 序号列居中
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, '楷体', 24)

# Add plan summary
add_body_paragraph('计划主线：数据标准先行 → 图像处理生成候选拓扑 → 分析引擎完成课程化推导 → 展示层打通手机拍照与电脑解题联动。', font_size=24, first_line_indent=420, line_spacing=348)

# （三）经费预算及使用计划
add_subsection_heading('（三）经费预算及使用计划')

add_numbered_item('1、调研、差旅费：500元')
add_body_paragraph('用于购买参考教材、收集典型题目和调研同类工具。')

add_numbered_item('2、用于项目研发的元器件、软硬件测试、小型硬件购置费等：1000元')
add_body_paragraph('用于购买测试用的电路板、元器件、摄像头等硬件设备，以及软件授权费用。')

add_numbered_item('3、资料购置、打印、复印、印刷等费用：300元')
add_body_paragraph('用于购买技术书籍、打印测试文档和项目报告。')

add_numbered_item('4、学生撰写与项目有关的论文版面费、申请专利费等：200元')
add_body_paragraph('用于项目成果的论文发表和专利申请。')

add_body_paragraph('经费使用计划：第一阶段（数据层建设）使用经费的30%，第二阶段（图像处理层开发）使用经费的30%，第三阶段（分析引擎层构建）使用经费的20%，第四阶段（展示层制作）使用经费的20%。', font_size=24, first_line_indent=420, line_spacing=348)

# ===== 五、项目预期成果 =====
add_section_heading('五、项目预期成果')

add_numbered_item('1、软件成果')
add_body_paragraph('项目预期形成一个本地运行的原型系统，包括以下核心组件：')
add_body_paragraph('（1）数据层：类型化电路图数据结构（TypeScript / Zod Schema）、元件库、Benchmark 数据集和 SQLite 任务队列，为系统提供统一的数据基础。')
add_body_paragraph('（2）图像处理层：基于 OpenCV 的预处理模块、元件检测 + OCR 参数提取、导线/节点/环路识别，生成带置信度的候选拓扑并提示可疑点。')
add_body_paragraph('（3）分析引擎层：课程规则层、符号工具层，输出结构化推导、关键方程、等效电路与可解释步骤。')
add_body_paragraph('（4）展示层：网页电路图编辑器、分析结果展示、LaTeX 公式渲染、报告导出、手机拍照与电脑解题联动。')

add_numbered_item('2、性能指标')
add_body_paragraph('第一阶段目标：在用户确认后的电路图上实现目标求解量正确率不低于 90%，LaTeX 渲染成功率不低于 95%，并完成可演示的端到端学习辅助流程。')

add_numbered_item('3、学术成果')
add_body_paragraph('（1）建立覆盖典型模电题型的开发测试集（选取 30 道典型教材题），为后续迭代和开源研究提供可复现依据。')
add_body_paragraph('（2）发表与项目相关的学术论文 1-2 篇。')
add_body_paragraph('（3）申请软件著作权 1 项。')

add_numbered_item('4、应用价值')
add_body_paragraph('项目完成后，可为学生提供一种可修正、可验证、可解释的电路学习工具，帮助其完成习题分析、理解等效模型、整理公式和撰写报告。后续可在现有架构基础上扩展电路分析基础题和数字电路分析，并逐步完善元件库、评测集和教学解释能力。')

# Save document
doc.save('files/开题报告-第四五部分.docx')
print('Document saved to files/开题报告-第四五部分.docx')
