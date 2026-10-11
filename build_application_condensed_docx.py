# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path("C:/Users/WZH/AppData/Local/Temp/circuit_project_application_condensed.docx")


def set_run_font(run, name="宋体", size=11, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def set_para(paragraph, first_line=True, before=0, after=6, line=1.45):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if first_line:
        fmt.first_line_indent = Cm(0.74)


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(text)
    set_run_font(r, "黑体", 20, True)


def add_meta(doc, label, value):
    p = doc.add_paragraph()
    set_para(p, first_line=False, after=3, line=1.25)
    r = p.add_run(label)
    set_run_font(r, "宋体", 11, True)
    for i, line in enumerate(value.split("\n")):
        if i:
            p.add_run("\n")
            rr = p.add_run(" " * 8 + line)
        else:
            rr = p.add_run(line)
        set_run_font(rr, "宋体", 11, False)


def add_heading(doc, text):
    p = doc.add_paragraph()
    set_para(p, first_line=False, before=10, after=6, line=1.2)
    r = p.add_run(text)
    set_run_font(r, "黑体", 13, True)


def add_body(doc, text):
    p = doc.add_paragraph()
    set_para(p)
    r = p.add_run(text)
    set_run_font(r)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.4)

styles = doc.styles
styles["Normal"].font.name = "宋体"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
styles["Normal"].font.size = Pt(11)

add_title(doc, "立项申请书")
add_meta(doc, "题目：", "基于视觉识别与符号推理的电路智能分析与教学辅助系统")
add_meta(doc, "学院：", "【请填写学院】")
add_meta(doc, "专业：", "【请填写专业】")
add_meta(doc, "学生姓名：", "【姓名1 学号】\n【姓名2 学号】\n【姓名3 学号】")

add_heading(doc, "申请原因")
add_body(
    doc,
    "电路分析、模拟电子技术和数字电子技术是电子信息类专业的重要基础课程，但其学习过程往往同时涉及电路图识别、拓扑建模、等效变换、公式推导和结果解释。尤其在二极管、三极管和 MOS 管基本电路中，学生不仅要看懂电路连接关系，还要完成直流等效、交流小信号等效、静态工作点、增益、输入电阻和输出电阻等分析，学习门槛较高。",
)
add_body(
    doc,
    "现有工具难以完整覆盖这一学习链路。专业仿真软件如 Multisim 偏重仿真和实验验证，不侧重规范解题步骤；CircuitMath 等工具能进行公式分析，但通常需要用户手动搭建电路；拍照搜题和通用大模型可以给出说明，却缺少可验证、可修正、可视化的电路拓扑和专业公式输出。因此，有必要研发一套面向课程学习的本地化电路智能分析与教学辅助系统。",
)

add_heading(doc, "项目简介")
add_body(
    doc,
    "本项目拟开发一套基于视觉识别与符号推理的电路智能分析与教学辅助系统。系统以电脑端本地软件为主体，支持手机在同一局域网内扫码拍照上传教材或习题册中的电路图；电脑端完成图像处理、候选拓扑生成、用户校正、自动分析和结果展示，所有数据和计算均在本地完成，不依赖云端服务。",
)
add_body(
    doc,
    "第一阶段重点面向模拟电子技术中的二极管、BJT 和 MOS 基本电路，建立 30 道典型教材题作为 benchmark。系统将生成可编辑的类型化电路图，用户确认节点、元件和参数后，系统自动生成直流等效电路图、交流小信号等效电路图、关键方程、目标参数、LaTeX 公式、步骤解释和知识点标注，并为后续扩展到电路分析和数字电路分析预留架构接口。",
)

add_heading(doc, "项目特色和创新点")
add_body(
    doc,
    "项目的核心创新在于将“电路图采集—拓扑校正—符号分析—等效电路展示—教学解释—LaTeX 导出”连成闭环。系统内部以类型化电路图作为唯一事实来源，SPICE / netlist 作为导出和校验层，避免前端画布、图像识别和分析引擎各自维护不同数据结构；同时采用轻量网页电路图编辑器和层级化元件库，兼顾第一阶段实现难度与后续扩展能力。",
)
add_body(
    doc,
    "分析引擎采用“课程规则层 + 符号/仿真工具层”的混合方案：课程规则层负责直流/交流等效、小信号模型和题型逻辑，SymPy、Lcapy、SPICE 用于符号求解、线性电路分析和数值校验。评测方面，项目将构建多层 ground truth，不只判断最终答案，还分别评测元件、节点、参数、等效电路、方程、LaTeX 和解释输出，从而具备开源研究和教学应用价值。",
)

add_heading(doc, "项目难点")
add_body(
    doc,
    "项目难点首先在于图像到拓扑的可靠转换。手机拍摄图片可能存在角度、光照、模糊、断线和符号风格差异，系统需要通过本地图像处理生成候选拓扑，并把低置信度元件、疑似断线和未连接端口等可疑点显式呈现给用户确认，避免错误拓扑直接进入分析流程。",
)
add_body(
    doc,
    "第二个难点在于分析结果的专业性和稳定性。模电题涉及工作区判断、等效电路生成、符号方程推导和表达式化简，若完全依赖大模型或单一工具，结果难以验证。项目拟采用插件式分析流程、分题型策略、符号计算超时保护和 SPICE 校验机制，确保结果可追溯、可解释、可评测。",
)

add_heading(doc, "项目预期成果")
add_body(
    doc,
    "项目预期形成一个本地运行的原型系统，包括电脑端网页应用、局域网手机采集页面、类型化电路图数据结构、轻量电路图编辑器、模电基础分析引擎、LaTeX 输出模块和 30 道典型题 benchmark。第一阶段目标是在用户确认后的电路图上实现目标求解量正确率不低于 90%，LaTeX 渲染成功率不低于 95%，并完成可演示的端到端学习辅助流程。",
)
add_body(
    doc,
    "项目完成后，可为学生提供一种可修正、可验证、可解释的电路学习工具，帮助其完成习题分析、理解等效模型、整理公式和撰写报告。后续可在现有架构基础上扩展电路分析基础题和数字电路分析，并逐步完善元件库、评测集和教学解释能力。恳请予以立项支持。",
)

doc.save(OUT)
print(OUT)
